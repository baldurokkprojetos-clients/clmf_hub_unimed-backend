"""
Service da OP2 ImprimirEvolucao (rotina 'clmf_imprimir_evolucao').

Responsabilidades:
  - parse da planilha modelo de evolucoes (openpyxl read_only)
  - agrupamento por (idPaciente, DataExec) -> 1 job por par, encadeando itens[guia/ID_prof](horas)
  - criacao dos jobs pendentes + itens PENDENTE em evolucao_itens
  - reprocesso integral no upload: pares que ja possuem job tem TODOS os itens
    do arquivo regravados como PENDENTE (lote novo) e jobs finalizados
    reenfileirados com flag 'reprocesso' — o arquivo subido e a sessao, nada
    do que ele contem e pulado
  - geracao do xlsx de status (compartilhado entre rota e script CLI)

Planilha modelo (aba ativa; cabecalho na primeira linha):
  idPaciente, nomePaciente, Guia, DataExec, Profissional, CodTerapia, Terapia,
  Sessoes, HoraInicial, Lote, status, tipo, ID_prof

Do JSON do job participam apenas: idPaciente, nomePaciente, Guia, DataExec, Terapia,
ID_prof, HoraInicial (+"login"/"senha" quando configurados no env do backend).
CodTerapia e Profissional NAO entram (profissao_id vem do select do portal pela
Terapia; nome do profissional e resolvido no select pelo ID_prof).
"""
import io
import os
import re
import threading
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from sqlalchemy import tuple_
from sqlalchemy.orm import Session

from models import Carteirinha, EvolucaoItem, Job

ROTINA_EVOLUCAO = "clmf_imprimir_evolucao"

# Carteirinha ancora criada pela migration 0032: o dispatcher exige carteirinha_id
# valido (dispatcher.py le job.carteirinha_rel.carteirinha). Evolucoes nao tem
# carteirinha no dominio — todos os jobs da rotina apontam para esta unica linha.
CARTEIRINHA_ANCORA = "EVOLUCOES-CLMF"

# Colunas obrigatorias da planilha (subset que segue no job)
REQUIRED_HEADERS = ["idPaciente", "nomePaciente", "Guia", "DataExec", "Terapia", "HoraInicial", "ID_prof"]

# Janela do filtro no portal: dataExec +/- 30 dias
JANELA_DIAS = 30

# Limites do relatorio de erros retornado pelo upload
MAX_ERROS_RETORNO = 50

# Criacao de jobs em LOTES: 1 commit a cada N jobs (commit-por-job via WAN ao
# Supabase custa ~1 job/s; com lote de 200, 10k jobs saem em ~1 min).
CRIACAO_BATCH_JOBS = 200

# Reconciliacao em LOTES: mesmo principio do CRIACAO_BATCH_JOBS — 1 commit a
# cada N pares reprocessados (UPDATE de itens + merge de params + reenqueue).
RECONCILIACAO_BATCH_PARES = 50


# ─── Parsing de celulas ─────────────────────────────────────────────────────

def _norm_header(h) -> str:
    return str(h or "").strip()


def _parse_data(v):
    """Valor de celula DataExec -> date | None."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_hora_hhmm(v) -> str | None:
    """'07:00:00' / '07:00' / time -> '0700' (formato do JSON do job)."""
    s = str(v or "").strip()
    m = re.match(r"^(\d{1,2})[:h]?(\d{2})", s)
    if m:
        hh, mm = int(m.group(1)), m.group(2)
        if 0 <= hh <= 23:
            return f"{hh:02d}{mm}"
    return None


# ─── Parse da planilha → payloads de job ────────────────────────────────────

def parse_planilha(contents: bytes) -> tuple[list[dict], list[str]]:
    """Retorna (jobs_payload, erros).

    jobs_payload: lista ordenada por (idPaciente, DataExec) — jobs do mesmo paciente
    ficam com ids consecutivos, favorecendo a estrategia de afinidade do worker.
    Cada payload segue o contrato de params do job (ver docstring do modulo).
    """
    wb = load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return [], ["Planilha vazia (sem cabecalho)."]

    header = {col: _norm_header(h) for col, h in enumerate(header_row) if h is not None}
    faltando = [c for c in REQUIRED_HEADERS if c not in header.values()]
    if faltando:
        raise ValueError(f"Cabecalho invalido. Colunas obrigatorias ausentes: {', '.join(faltando)}. "
                         f"Esperado: {', '.join(REQUIRED_HEADERS)}")

    col = {name: idx for idx, name in header.items()}
    erros: list[str] = []
    grupos: dict[tuple[int, date], dict] = {}
    nome_por_paciente: dict[int, str] = {}

    for lin, row in enumerate(rows, start=2):
        if row is None or all(c is None for c in row):
            continue

        def cell(name):
            idx = col.get(name)
            return row[idx] if idx is not None and idx < len(row) else None

        try:
            id_paciente = int(str(cell("idPaciente")).strip())
        except (TypeError, ValueError):
            erros.append(f"Linha {lin}: idPaciente vazio/invalido ('{cell('idPaciente')}')")
            continue
        data_exec = _parse_data(cell("DataExec"))
        if data_exec is None:
            erros.append(f"Linha {lin}: DataExec vazio/invalido ('{cell('DataExec')}')")
            continue
        guia = str(cell("Guia") or "").strip()
        if not guia:
            erros.append(f"Linha {lin}: Guia vazia")
            continue
        terapia = str(cell("Terapia") or "").strip()
        if not terapia:
            erros.append(f"Linha {lin}: Terapia vazia")
            continue
        try:
            id_prof = int(str(cell("ID_prof")).strip())
        except (TypeError, ValueError):
            erros.append(f"Linha {lin}: ID_prof vazio/invalido ('{cell('ID_prof')}')")
            continue
        hora = _parse_hora_hhmm(cell("HoraInicial"))
        if hora is None:
            erros.append(f"Linha {lin}: HoraInicial vazio/invalido ('{cell('HoraInicial')}')")
            continue

        nome = str(cell("nomePaciente") or "").strip()
        if id_paciente not in nome_por_paciente and nome:
            nome_por_paciente[id_paciente] = nome

        chave = (id_paciente, data_exec)
        grupo = grupos.setdefault(chave, {
            "idPaciente": id_paciente,
            "dataExec": data_exec.isoformat(),
            "itens": {},  # (guia, id_prof) -> {"guia","ID_prof","terapia","horas":[]}
        })
        item = grupo["itens"].setdefault((guia, id_prof), {
            "guia": guia, "ID_prof": id_prof, "terapia": terapia, "horas": []
        })
        if hora not in item["horas"]:
            item["horas"].append(hora)

    payloads = []
    for (id_paciente, data_exec), grupo in sorted(grupos.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        data_ini = (data_exec - timedelta(days=JANELA_DIAS)).isoformat()
        data_fim = (data_exec + timedelta(days=JANELA_DIAS)).isoformat()
        payloads.append({
            "idPaciente": id_paciente,
            "nomePaciente": nome_por_paciente.get(id_paciente, ""),
            "dataExec": data_exec.isoformat(),
            "dataInicial": data_ini,
            "dataFinal": data_fim,
            "itens": [it for it in grupo["itens"].values()],
        })
    return payloads, erros


# ─── Criacao dos jobs ───────────────────────────────────────────────────────

def _credenciais_params() -> dict:
    """login/senha entram em params (spec) quando configurados no env do backend;
    ausentes, o scraper usa o proprio env do worker (CLMF_LOGIN/CLMF_PASSWORD)."""
    login = os.getenv("CLMF_LOGIN", "").strip()
    senha = os.getenv("CLMF_PASSWORD", "").strip()
    out = {}
    if login:
        out["login"] = login
    if senha:
        out["senha"] = senha
    return out


def nome_lote(filename: str) -> str:
    return f"{filename}:{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"


def _pares_existentes(db: Session) -> set[tuple[str, str]]:
    """(idPaciente, dataExec) que JÁ possuem job da rotina (qualquer status).

    Projeção leve (2 colunas) para a estimativa do resumo do upload: pares
    listados aqui serão REPROCESSADOS integralmente pelo upload — o arquivo
    subido é a sessão e nenhum item dele é pulado.
    """
    rows = db.query(
        Job.params["idPaciente"].astext,
        Job.params["dataExec"].astext,
    ).filter(Job.rotina == ROTINA_EVOLUCAO).all()
    return {(r[0], r[1]) for r in rows if r[0] is not None}


def _jobs_por_par(db: Session) -> dict[tuple[str, str], Job]:
    """Jobs da rotina indexados por (idPaciente, dataExec) — qualquer status.

    Usado pela reconciliação: o reprocesso acontece SEMPRE no job existente
    do par (nunca cria segundo job — índice único uq_jobs_evolucao_par).
    """
    jobs = db.query(Job).filter(Job.rotina == ROTINA_EVOLUCAO).all()
    return {(str(j.params["idPaciente"]), j.params["dataExec"]): j
            for j in jobs if (j.params or {}).get("idPaciente") is not None}


def _itens_por_job(db: Session, job_ids: list[int]) -> dict[int, dict[tuple[str, int, str], str]]:
    """Itens dos jobs indexados por job_id → {(guia, profissional_id, hora): status}."""
    if not job_ids:
        return {}
    rows = db.query(
        EvolucaoItem.job_id, EvolucaoItem.guia, EvolucaoItem.profissional_id,
        EvolucaoItem.hora_inicial, EvolucaoItem.status,
    ).filter(EvolucaoItem.job_id.in_(job_ids)).all()
    out: dict[int, dict[tuple[str, int, str], str]] = {}
    for jid, guia, prof, hora, status in rows:
        out.setdefault(jid, {})[(guia, prof, hora)] = status
    return out


def _commit_lote_jobs(db: Session, payloads: list[dict], lote: str, ancora_id: int) -> int:
    """Insere um lote de jobs + itens PENDENTE em UMA transação (flush unico
    popula os job.id). Retorna nº de itens criados."""
    job_objs = []
    for payload in payloads:
        job = Job(
            carteirinha_id=ancora_id,
            rotina=ROTINA_EVOLUCAO,
            params=payload,  # credenciais já embutidas
            status="pending",
        )
        db.add(job)
        job_objs.append(job)
    db.flush()  # popula os ids dos jobs

    itens = []
    for job_obj, payload in zip(job_objs, payloads):
        itens.extend(_itens_do_payload(job_obj.id, payload, lote))
    db.add_all(itens)
    db.commit()
    return len(itens)


def _itens_do_payload(job_id: int, payload: dict, lote: str) -> list[EvolucaoItem]:
    out = []
    for item in payload["itens"]:
        for hora in item["horas"]:
            out.append(EvolucaoItem(
                job_id=job_id,
                lote=lote,
                id_paciente=payload["idPaciente"],
                nome_paciente=payload["nomePaciente"],
                guia=item["guia"],
                data_exec=date.fromisoformat(payload["dataExec"]),
                profissional_id=item["ID_prof"],
                terapia=item["terapia"],
                hora_inicial=hora,
                status="PENDENTE",
            ))
    return out


# ─── Reprocesso integral de pares já importados (upload = sessão) ───────────

def _reconciliar_par(db: Session, payload: dict, job: Job,
                     existentes: dict[tuple[str, int, str], str],
                     lote: str, incremento_min: int) -> dict:
    """Reprocessa integralmente, NO JOB EXISTENTE do par, um payload do arquivo.

    (1) itens FALTANTES (no arquivo, não no job): merge em params["itens"] (o
    worker processa params, não a tabela) + insert de EvolucaoItem PENDENTE;
    (2) itens PRESENTES: TODO item do arquivo volta a PENDENTE com o lote novo
    (status OK/ERRO/PENDENTE, motivo limpo) — a retomada do worker só pula
    grupos 100% OK, logo todos são refeitos e os PDFs rebaixados;
    (3) job finalizado (success/error) é reenfileirado com a flag 'reprocesso'
    (a seleção do worker ignora claims antigos e habilita candidatos já
    preenchidos — sem isso, os candidatos do run anterior continuam reservados
    em evolucao_claims e os itens viram "candidatos insuficientes").

    Job em 'processing': apenas (1) — a run em curso leu params no início e o
    reset disputaria estado com o worker; o reprocesso fica adiado.
    """
    horas_por_grupo: dict[tuple[str, int], list[str]] = {}
    for it in payload["itens"]:
        hs = horas_por_grupo.setdefault((it["guia"], it["ID_prof"]), [])
        for h in it["horas"]:
            if h not in hs:
                hs.append(h)
    chaves_arquivo = {(g, p, h) for (g, p), hs in horas_por_grupo.items() for h in hs}
    faltantes = chaves_arquivo - existentes.keys()
    a_resetar = chaves_arquivo & existentes.keys()

    stats = {"pares_reprocessados": 0, "itens_resetados": 0, "itens_novos": 0,
             "jobs_reenfileirados": 0, "jobs_adiados": 0}

    params = dict(job.params or {})
    itens_params = [dict(i) for i in params.get("itens", [])]
    index = {(i["guia"], i["ID_prof"]): i for i in itens_params}

    if faltantes:
        data_exec = date.fromisoformat(payload["dataExec"])
        falt_por_grupo: dict[tuple[str, int], list[str]] = {}
        for g, p, h in faltantes:
            falt_por_grupo.setdefault((g, p), []).append(h)
        novos_objs = []
        for (guia, id_prof), horas in falt_por_grupo.items():
            fonte = next(i for i in payload["itens"]
                         if i["guia"] == guia and i["ID_prof"] == id_prof)
            if (guia, id_prof) in index:
                for h in horas:
                    if h not in index[(guia, id_prof)]["horas"]:
                        index[(guia, id_prof)]["horas"].append(h)
                terapia_item = index[(guia, id_prof)].get("terapia", "")
            else:
                itens_params.append({"guia": guia, "ID_prof": id_prof,
                                     "terapia": fonte["terapia"], "horas": list(horas)})
                terapia_item = fonte["terapia"]
            for h in horas:
                novos_objs.append(EvolucaoItem(
                    job_id=job.id, lote=lote, id_paciente=payload["idPaciente"],
                    nome_paciente=payload.get("nomePaciente") or "", guia=guia,
                    data_exec=data_exec, profissional_id=id_prof,
                    terapia=terapia_item, hora_inicial=h, status="PENDENTE",
                ))
        params["itens"] = itens_params
        db.add_all(novos_objs)
        stats["itens_novos"] = len(novos_objs)

    params["incrementoMin"] = incremento_min

    if job.status == "processing":
        job.params = params  # reatribuição p/ SQLAlchemy detectar mudança no JSONB
        stats["jobs_adiados"] = 1
        return stats

    if a_resetar:
        db.query(EvolucaoItem).filter(
            EvolucaoItem.job_id == job.id,
            tuple_(EvolucaoItem.guia, EvolucaoItem.profissional_id,
                   EvolucaoItem.hora_inicial).in_(sorted(a_resetar)),
        ).update({"status": "PENDENTE", "motivo": None, "lote": lote,
                  "updated_at": datetime.utcnow()}, synchronize_session=False)
        stats["itens_resetados"] = len(a_resetar)

    if job.status in ("success", "error"):
        params["reprocesso"] = True
        job.status = "pending"
        job.attempts = 0
        job.locked_by = None
        job.updated_at = datetime.utcnow()
        stats["jobs_reenfileirados"] = 1
    job.params = params  # reatribuição p/ SQLAlchemy detectar mudança no JSONB
    stats["pares_reprocessados"] = 1
    return stats


def _reconciliar_pares(db: Session, pares: list[tuple[dict, Job]], lote: str,
                       incremento_min: int) -> tuple[dict, list[str]]:
    """Driver do reprocesso: 1 commit a cada RECONCILIACAO_BATCH_PARES pares,
    com isolamento de falha por lote (mesma rationale de _commit_lote_jobs).
    Retorna (stats, falhas)."""
    itens_status = _itens_por_job(db, [job.id for _, job in pares])

    stats = {"pares_reprocessados": 0, "itens_resetados": 0, "itens_novos": 0,
             "jobs_reenfileirados": 0, "jobs_adiados": 0}
    falhas: list[str] = []
    em_lote: list[tuple[dict, dict]] = []

    def _commit_lote():
        try:
            db.commit()
            for _, res in em_lote:
                for k, v in res.items():
                    stats[k] += v
        except Exception as e:
            db.rollback()
            primeiro = em_lote[0][0]
            falhas.append(f"lote de {len(em_lote)} pares a partir de "
                          f"{primeiro['idPaciente']}/{primeiro['dataExec']}: {e}")
        em_lote.clear()

    for payload, job in pares:
        try:
            res = _reconciliar_par(db, payload, job, itens_status.get(job.id, {}),
                                   lote, incremento_min)
        except Exception as e:
            # rollback descarta também pares acumulados sem commit no lote corrente
            db.rollback()
            em_lote.clear()
            falhas.append(f"par {payload['idPaciente']}/{payload['dataExec']}: {e}")
            continue
        em_lote.append((payload, res))
        if len(em_lote) >= RECONCILIACAO_BATCH_PARES:
            _commit_lote()
    if em_lote:
        _commit_lote()
    return stats, falhas


def criar_jobs_evolucao(db: Session, payloads: list[dict], filename: str, lote: str | None = None,
                        incremento_min: int = 60) -> dict:
    """O upload é uma SESSÃO COMPLETA sobre o arquivo enviado:

      - pares INÉDITOS (idPaciente, DataExec) → cria 1 job novo + itens
        PENDENTE em evolucao_itens (commits em lote de CRIACAO_BATCH_JOBS,
        com isolamento de falha por lote);
      - pares que JÁ possuem job (qualquer lote/status) → reprocesso integral
        via _reconciliar_pares: todo item do arquivo volta a PENDENTE com o
        lote novo, itens faltantes são acrescentados ao job e jobs
        finalizados são reenfileirados com a flag 'reprocesso'.

    incremento_min: duração da sessão — horaFinal = horaInicial + incremento
    (60 = 1h, padrão/histórico; 30 = 30min). Vai no params do job como
    'incrementoMin'; jobs sem o campo continuam em 60 (retrocompatível). Pares
    reprocessados também são reflagados com o incremento vigente do upload.
    """
    ancora = db.query(Carteirinha).filter(Carteirinha.carteirinha == CARTEIRINHA_ANCORA).first()
    if not ancora:
        raise RuntimeError(f"Carteirinha ancora '{CARTEIRINHA_ANCORA}' ausente — rode a migration 0032.")

    lote = lote or nome_lote(filename)
    credenciais = _credenciais_params()
    jobs_por_par = _jobs_por_par(db)

    jobs_ok, itens_ok, falhas = 0, 0, []
    pendentes: list[dict] = []
    reconciliar: list[tuple[dict, Job]] = []
    vistos: set[tuple[str, str]] = set()  # defesa extra: o parse já agrupa por par único

    for payload in payloads:
        chave = (str(payload["idPaciente"]), payload["dataExec"])
        if chave in vistos:
            continue
        vistos.add(chave)
        job = jobs_por_par.get(chave)
        if job is not None:
            reconciliar.append((payload, job))  # reprocesso integral no job existente
            continue
        pendentes.append({**payload, "incrementoMin": incremento_min, **credenciais})

        if len(pendentes) >= CRIACAO_BATCH_JOBS:
            try:
                itens_ok += _commit_lote_jobs(db, pendentes, lote, ancora.id)
                jobs_ok += len(pendentes)
            except Exception as e:
                db.rollback()
                falhas.append(f"lote de {len(pendentes)} jobs a partir de "
                              f"{pendentes[0]['idPaciente']}/{pendentes[0]['dataExec']}: {e}")
            pendentes = []

    if pendentes:
        try:
            itens_ok += _commit_lote_jobs(db, pendentes, lote, ancora.id)
            jobs_ok += len(pendentes)
        except Exception as e:
            db.rollback()
            falhas.append(f"lote final de {len(pendentes)} jobs: {e}")

    resumo = {"lote": lote, "jobs": jobs_ok, "itens": itens_ok,
              "incrementoMin": incremento_min}
    if reconciliar:
        stats, falhas_rec = _reconciliar_pares(db, reconciliar, lote, incremento_min)
        resumo.update(stats)
        falhas.extend(falhas_rec)
    else:
        resumo.update({"pares_reprocessados": 0, "itens_resetados": 0,
                       "itens_novos": 0, "jobs_reenfileirados": 0, "jobs_adiados": 0})
    resumo["falhas"] = falhas[:MAX_ERROS_RETORNO]
    return resumo


def criar_jobs_background(payloads: list[dict], filename: str, lote: str, incremento_min: int = 60):
    """Executa a criação/reprocesso em thread própria (session dedicada).

    Upload de planilhas grandes (20k+ linhas) não pode travar o event loop nem
    estourar timeout do proxy (Render): o endpoint responde imediatamente e o
    trabalho corre aqui; o progresso aparece no painel via /evolucoes/jobs.
    """
    try:
        from database import SessionLocal
        db = SessionLocal()
        try:
            resumo = criar_jobs_evolucao(db, payloads, filename, lote, incremento_min=incremento_min)
            print(f"[evolucoes] lote {lote}: {resumo['jobs']} novos jobs / {resumo['itens']} itens criados "
                  f"[incremento {resumo.get('incrementoMin', 60)}min]; "
                  f"reprocesso: {resumo['pares_reprocessados']} pares "
                  f"({resumo['itens_resetados']} itens resetados, {resumo['itens_novos']} itens novos, "
                  f"{resumo['jobs_reenfileirados']} jobs reenfileirados, {resumo['jobs_adiados']} adiados), "
                  f"falhas={len(resumo['falhas'])}")
        finally:
            db.close()
    except Exception as e:
        print(f"[evolucoes] ERRO na criacao do lote {lote}: {e}")


# ─── Export de status (xlsx) ────────────────────────────────────────────────

EXPORT_HEADERS = ["idPaciente", "nomePaciente", "guia", "data", "profissional_id",
                  "terapia", "profissao_id", "horaInicial", "status", "motivo", "pdfPath", "lote"]


def gerar_xlsx_status(db: Session, lote: str | None = None, job_id: int | None = None,
                      status: str | None = None) -> io.BytesIO:
    """Gera o xlsx de status a partir de evolucao_itens (padrao routes/guias.py)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Evolucoes")
    ws.append(EXPORT_HEADERS)

    query = db.query(
        EvolucaoItem.id_paciente, EvolucaoItem.nome_paciente, EvolucaoItem.guia,
        EvolucaoItem.data_exec, EvolucaoItem.profissional_id, EvolucaoItem.terapia,
        EvolucaoItem.profissao_id, EvolucaoItem.hora_inicial, EvolucaoItem.status,
        EvolucaoItem.motivo, EvolucaoItem.pdf_path, EvolucaoItem.lote,
    )
    if lote:
        query = query.filter(EvolucaoItem.lote == lote)
    if job_id:
        query = query.filter(EvolucaoItem.job_id == job_id)
    if status:
        query = query.filter(EvolucaoItem.status == status)

    for row in query.yield_per(1000).order_by(EvolucaoItem.id_paciente, EvolucaoItem.data_exec,
                                              EvolucaoItem.guia, EvolucaoItem.hora_inicial):
        ws.append([
            row.id_paciente, row.nome_paciente or "", row.guia or "",
            row.data_exec.strftime("%d/%m/%Y") if row.data_exec else "",
            row.profissional_id, row.terapia or "", row.profissao_id or "",
            row.hora_inicial or "", row.status, row.motivo or "",
            row.pdf_path or "", row.lote or "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
